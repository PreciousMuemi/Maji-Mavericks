import io
from zipfile import ZipFile

import fitz
import pandas as pd
import pytest
from docx import Document

from ai.config import Settings
from ai.document_parser import DocumentError, DocumentParser


def parser(**kwargs):
    return DocumentParser(Settings(_env_file=None, **kwargs))


def test_csv_preserves_identifiers_and_locators():
    result = parser().parse(b'id,value\n001,250\n', 'portfolio.csv', 'doc')
    assert result.segments[0].locator == 'sheet:csv:row:2'
    assert '001' in result.segments[0].text


def test_xlsx_multiple_sheets():
    stream = io.BytesIO()
    with pd.ExcelWriter(stream, engine='openpyxl') as writer:
        pd.DataFrame({'id': ['001']}).to_excel(writer, sheet_name='Assets', index=False)
        pd.DataFrame({'claim': ['C1']}).to_excel(writer, sheet_name='Claims', index=False)
    result = parser().parse(stream.getvalue(), 'book.xlsx', 'doc')
    assert [s.locator for s in result.segments] == ['sheet:Assets:row:1', 'sheet:Assets:row:2', 'sheet:Claims:row:1', 'sheet:Claims:row:2']


def test_docx_paragraph_and_table():
    document = Document()
    document.add_paragraph('Nzoia warehouse')
    document.add_table(rows=1, cols=1).cell(0, 0).text = 'KES 100'
    stream = io.BytesIO()
    document.save(stream)
    result = parser().parse(stream.getvalue(), 'policy.docx', 'doc')
    assert [s.locator for s in result.segments] == ['paragraph:1', 'table:1:row:1']


def test_pdf_text_and_scanned_page_warning():
    with fitz.open() as document:
        document.new_page().insert_text((72, 72), 'Warehouse insurance')
        document.new_page()
        content = document.tobytes()
    result = parser().parse(content, 'policy.pdf', 'doc')
    assert result.segments[0].locator == 'page:1'
    assert 'Warehouse insurance' in result.segments[0].text
    assert 'OCR is not implemented' in result.warnings[0]


@pytest.mark.parametrize('content,filename', [(b'', 'x.csv'), (b'broken', 'x.pdf'), (b'x', 'x.exe'), (b'broken', 'x.docx')])
def test_invalid_documents(content, filename):
    with pytest.raises(DocumentError):
        parser().parse(content, filename, 'doc')


def test_size_and_text_limits():
    with pytest.raises(DocumentError):
        parser(max_upload_bytes=2).parse(b'id\n001', 'x.csv', 'doc')
    with pytest.raises(DocumentError):
        parser(max_document_characters=2).parse(b'id\n001', 'x.csv', 'doc')


def test_archive_expansion_limit():
    stream = io.BytesIO()
    with ZipFile(stream, 'w') as archive:
        archive.writestr('large.xml', 'a' * 100)
    with pytest.raises(DocumentError, match='uncompressed'):
        parser(max_archive_uncompressed_bytes=50).parse(stream.getvalue(), 'x.xlsx', 'doc')


def test_encrypted_pdf_is_rejected():
    with fitz.open() as document:
        document.new_page()
        content = document.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw='owner', user_pw='reader')
    with pytest.raises(DocumentError, match='Encrypted'):
        parser().parse(content, 'encrypted.pdf', 'doc')


def test_excel_preserves_narrative_and_irregular_headings():
    stream = io.BytesIO()
    frame = pd.DataFrame([['CLIENT: Example Mill', ''], ['Revised asset list', ''], ['Structure', 'Area'], ['Warehouse', '120 m2']])
    with pd.ExcelWriter(stream, engine='openpyxl') as writer:
        frame.to_excel(writer, sheet_name='Offer', index=False, header=False)
    document = parser().parse(stream.getvalue(), 'offer.xlsx', 'doc')
    assert len(document.segments) == 4
    assert 'CLIENT: Example Mill' in document.segments[0].text
    assert all(segment.kind == 'table' for segment in document.segments)


def test_pdf_table_cells_and_page_numbers():
    with fitz.open() as document:
        page = document.new_page()
        for x in [50, 250, 450]:
            page.draw_line((x, 50), (x, 170))
        for y in [50, 90, 130, 170]:
            page.draw_line((50, y), (450, y))
        page.insert_text((60, 75), 'Building')
        page.insert_text((260, 75), 'Area m2')
        page.insert_text((60, 115), 'Warehouse')
        page.insert_text((260, 115), '120')
        page.insert_text((60, 155), 'Office')
        page.insert_text((260, 155), '80')
        content = document.tobytes()
    result = parser().parse(content, 'table.pdf', 'doc')
    tables = [s for s in result.segments if s.kind == 'table']
    assert tables and 'Warehouse' in tables[0].text and '120' in tables[0].text
    assert tables[0].page_number == 1
    assert tables[0].locator == 'page:1:table:1'
