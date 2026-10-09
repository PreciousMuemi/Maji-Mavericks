"""Bounded ingestion; archive contents are never extracted onto the filesystem."""
import csv
import io
import json
from pathlib import Path
from threading import RLock
from zipfile import ZipFile

import pymupdf as fitz
import pandas as pd
from docx import Document

from .config import Settings
from .schemas import DocumentSegment, ParsedDocument


class _DiscardDiagnostics:
    """Vendor exception text is untrusted and must not enter logs or JSON stdout."""
    def write(self, message):
        return len(message)

    def flush(self):
        pass


if hasattr(fitz, 'no_recommend_layout'):
    fitz.no_recommend_layout()
if hasattr(fitz, 'set_messages'):
    fitz.set_messages(stream=_DiscardDiagnostics())
fitz.TOOLS.mupdf_display_errors(False)
fitz.TOOLS.mupdf_display_warnings(False)
_PDF_LOCK = RLock()  # PyMuPDF table detection uses shared library state.


class DocumentError(ValueError):
    pass


class DocumentParser:
    supported_suffixes = {'.pdf', '.xlsx', '.csv', '.docx'}

    def __init__(self, settings: Settings):
        self.settings = settings

    def parse_file(self, file_path: str | Path, document_id: str) -> ParsedDocument:
        path = Path(file_path)
        if path.suffix.lower() not in self.supported_suffixes:
            raise DocumentError('Supported formats are PDF, XLSX, CSV, and DOCX')
        try:
            if not path.is_file():
                raise DocumentError('Document path must be a regular file')
            with path.open('rb') as stream:
                content = stream.read(self.settings.max_upload_bytes + 1)
        except OSError as exc:
            raise DocumentError('Document could not be read') from exc
        return self.parse(content, path.name, document_id)

    def parse(self, content: bytes, filename: str, document_id: str) -> ParsedDocument:
        if not content or len(content) > self.settings.max_upload_bytes:
            raise DocumentError('Document is empty or exceeds the upload limit')
        suffix = Path(filename).suffix.lower()
        if suffix not in self.supported_suffixes:
            raise DocumentError('Supported formats are PDF, XLSX, CSV, and DOCX')
        segments, warnings = [], []
        character_count = 0

        def append(locator, text, page=None, kind='text'):
            nonlocal character_count
            if not text.strip():
                return
            character_count += len(text)
            if character_count > self.settings.max_document_characters:
                raise DocumentError('Document text exceeds the extraction limit; split the document')
            segments.append(DocumentSegment(locator=locator, text=text, page_number=page, kind=kind))

        try:
            if suffix in {'.xlsx', '.docx'}:
                with ZipFile(io.BytesIO(content)) as archive:
                    entries = archive.infolist()
                    if sum(entry.file_size for entry in entries) > self.settings.max_archive_uncompressed_bytes or len(entries) > 10_000:
                        raise DocumentError('Archive exceeds the uncompressed size limit')
                    names = set(archive.namelist())
                    required = 'xl/workbook.xml' if suffix == '.xlsx' else 'word/document.xml'
                    if required not in names or '[Content_Types].xml' not in names:
                        raise DocumentError('File contents do not match the declared document type')
            if suffix == '.pdf':
                if not content[:1024].lstrip().startswith(b'%PDF-'):
                    raise DocumentError('File contents do not match PDF format')
                with _PDF_LOCK, fitz.open(stream=content, filetype='pdf') as pdf:
                    if pdf.needs_pass:
                        raise DocumentError('Encrypted PDFs are unsupported')
                    for i, page in enumerate(pdf):
                        number = i + 1
                        text = page.get_text(sort=True).strip()
                        if text:
                            append(f'page:{number}', text, number)
                        else:
                            warnings.append(f'Page {number} has no extractable text; OCR is not implemented')
                        try:
                            tables = [table.extract() for table in page.find_tables().tables]
                        except Exception:
                            warnings.append(f'Page {number} table detection was incomplete; review the page text')
                            tables = []
                        for j, table in enumerate(tables, 1):
                            append(f'page:{number}:table:{j}', json.dumps(table, ensure_ascii=False), number, 'table')
            elif suffix == '.xlsx':
                # header=None retains titles, merged headings and repeated headers.
                sheets = pd.read_excel(io.BytesIO(content), sheet_name=None, header=None, dtype=str, keep_default_na=False, engine='openpyxl')
                for sheet, frame in sheets.items():
                    for index, row in frame.fillna('').iterrows():
                        cells = row.tolist()
                        if any(cell.strip() for cell in cells):
                            append(f'sheet:{sheet}:row:{index + 1}', json.dumps(cells, ensure_ascii=False), kind='table')
                warnings.append('Spreadsheet cells use cached values; formulas are not executed and missing formula caches require review')
            elif suffix == '.csv':
                if b'\x00' in content:
                    raise DocumentError('CSV must be UTF-8 text')
                text = content.decode('utf-8-sig')
                try:
                    separator = csv.Sniffer().sniff(text[:8192], delimiters=',;\t|').delimiter
                except csv.Error:
                    separator = ','
                # pandas may infer an index from an extra leading cell or rename
                # duplicate headers. Either silently changes exposure evidence.
                rows = csv.reader(io.StringIO(text), delimiter=separator)
                header = next(rows, [])
                if len(set(header)) != len(header):
                    raise DocumentError('CSV headers must be unique; clarify ambiguous source columns')
                if any(row and len(row) != len(header) for row in rows):
                    raise DocumentError('CSV row widths do not match the declared columns')
                frame = pd.read_csv(io.StringIO(text), dtype=str, keep_default_na=False, sep=separator)
                for index, row in frame.fillna('').iterrows():
                    append(f'sheet:csv:row:{index + 2}', json.dumps(row.to_dict(), ensure_ascii=False), kind='table')
            elif suffix == '.docx':
                doc = Document(io.BytesIO(content))
                for i, paragraph in enumerate(doc.paragraphs, 1):
                    append(f'paragraph:{i}', paragraph.text)
                for i, table in enumerate(doc.tables, 1):
                    for j, row in enumerate(table.rows, 1):
                        append(f'table:{i}:row:{j}', json.dumps([cell.text for cell in row.cells], ensure_ascii=False), kind='table')
                warnings.append('DOCX pagination is not reliable without a layout renderer; paragraph and table references are preserved')
        except DocumentError:
            raise
        except Exception as exc:
            raise DocumentError('Document could not be parsed') from exc
        if not segments:
            warnings.append('No extractable text was found')
        return ParsedDocument(document_id=document_id, segments=segments, warnings=warnings)
